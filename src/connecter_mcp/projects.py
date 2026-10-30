"""Project output: FF&E schedules, presentation boards, client selection sheets and usage analytics."""

from __future__ import annotations

import base64
import html
import io
import json
import os
import re
import time
from collections import Counter, defaultdict
from pathlib import Path

from . import maxfile
from .ai import _brands_in, _first_previews, _materials_in, asset_image, describe
from .client import ConnecterError
from .config import HOME_DIR
from .maxtools import library_models_for, measurements
from .server import GUID_RE, _db, _get_asset, _texture_status, search_assets, tool

EXPORTS = HOME_DIR / "exports"
USAGE_CACHE = HOME_DIR / "usage_cache.json"
CODE_PREFIX = {"FURNITURE": "FUR", "DECOR": "DEC", "BATHROOM": "BTH", "KITCHEN": "KIT", "LIGHT": "LGT",
               "TREE": "PLT", "PLANT": "PLT", "MISC": "MSC"}


def _rows(assets: list[str] | None, tag: str, scene: str) -> list[dict]:
    """The assets to put in a document, with tags, from a list, a tag or the library models a scene uses."""
    with _db() as db:
        catalogue = {a.path.lower(): a for a in db.assets()}
    if scene:
        f = Path(scene)
        paths = [_texture_status(t, f)["resolved"] for t in maxfile.max_texture_paths(f)]
        found = library_models_for([p for p in paths if p])
        return [{"path": h["asset"], "tags": h["tags"], "description": catalogue[h["asset"].lower()].description} for h in found]
    if assets:
        out = []
        for a in assets:
            p = _get_asset(a)["FileFullPath"] if GUID_RE.match(a.strip()) else a
            c = catalogue.get(p.lower())
            out.append({"path": c.path if c else p, "tags": c.tags if c else [], "description": c.description if c else None})
        return out
    if tag:
        return search_assets(tag=tag, limit=100000)["assets"]
    raise ConnecterError("Give assets, a tag or a scene")


def _item(row: dict, previews: dict[str, str]) -> dict:
    p = Path(row["path"])
    tag = row["tags"][0] if row.get("tags") else ""
    text = f"{p.parent.name} {p.stem} {row.get('description') or ''}"
    dims = measurements().get(str(p).lower())
    return {"name": re.sub(r"\s+", " ", re.sub(r"[_]+", " ", p.stem)).strip(), "path": str(p), "tag": tag,
            "category": tag.split("/")[-1] if tag else p.parent.parent.name,
            "brand": (_brands_in(text) or [""])[0], "materials": ", ".join(_materials_in(text)[:4]),
            "description": row.get("description") or describe(str(p)),
            "dimensions": f"{dims['width_cm']:.0f} x {dims['depth_cm']:.0f} x {dims['height_cm']:.0f} cm" if dims else "",
            "image": asset_image(str(p), previews)}


def _prefix(tag: str) -> str:
    top = re.sub(r"^\d+_", "", tag.split("/")[0]).upper() if tag else ""
    return next((v for k, v in CODE_PREFIX.items() if k in top), (top[:3] or "ITM"))


def _thumb_bytes(path: str | None, size: int = 400) -> bytes | None:
    if not path or not Path(path).exists():
        return None
    from PIL import Image

    try:
        with Image.open(path) as im:
            im = im.convert("RGB")
            im.thumbnail((size, size))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=82)
            return buf.getvalue()
    except Exception:
        return None


def _out(output: str, name: str, ext: str) -> Path:
    if output:
        p = Path(output)
        return p if p.suffix else p / f"{name}{ext}"
    EXPORTS.mkdir(parents=True, exist_ok=True)
    return EXPORTS / f"{re.sub(r'[^A-Za-z0-9_-]+', '_', name)}_{time.strftime('%Y%m%d-%H%M')}{ext}"


@tool
def ffe_schedule(assets: list[str] | None = None, tag: str = "", scene: str = "", title: str = "FF&E Schedule",
                 quantities: dict[str, int] | None = None, output: str = "", include_images: bool = True) -> dict:
    """An FF&E schedule (Excel, plus CSV) for a project: item codes by category (FUR-01, LGT-03 ...), name,
    category, brand, description, materials, dimensions (from measure_asset), quantity, image and file.
    Items come from a list of assets, a tag, or the library models used by a .max scene."""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    previews = _first_previews()
    items = [_item(r, previews) for r in _rows(assets, tag, scene)]
    if not items:
        raise ConnecterError("No items found")
    items.sort(key=lambda i: (i["tag"], i["name"].lower()))
    counters: Counter = Counter()
    for i in items:
        pre = _prefix(i["tag"])
        counters[pre] += 1
        i["code"] = f"{pre}-{counters[pre]:02d}"
        i["qty"] = (quantities or {}).get(i["path"], (quantities or {}).get(i["name"], 1))
    path = _out(output, title, ".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = ["Code", "Image", "Item", "Category", "Brand", "Description", "Materials", "Dimensions", "Qty", "File"]
    wb = Workbook()
    ws = wb.active
    ws.title = "FF&E"
    ws.append([title])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([f"Generated {time.strftime('%Y-%m-%d')}"])
    ws.append(cols)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
    widths = [9, 18, 30, 18, 16, 50, 22, 20, 6, 60]
    for idx, w in enumerate(widths):
        ws.column_dimensions[chr(65 + idx)].width = w
    for n, i in enumerate(items, start=4):
        ws.append([i["code"], "", i["name"], i["category"], i["brand"], i["description"], i["materials"], i["dimensions"], i["qty"], i["path"]])
        for c in ws[n]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if include_images and (data := _thumb_bytes(i["image"], 120)):
            img = XLImage(io.BytesIO(data))
            ws.add_image(img, f"B{n}")
            ws.row_dimensions[n].height = 95
    ws.freeze_panes = "A4"
    wb.save(path)
    csv_path = path.with_suffix(".csv")
    import csv

    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c for c in cols if c != "Image"])
        for i in items:
            w.writerow([i["code"], i["name"], i["category"], i["brand"], i["description"], i["materials"], i["dimensions"], i["qty"], i["path"]])
    return {"xlsx": str(path), "csv": str(csv_path), "items": len(items),
            "missing_dimensions": sum(1 for i in items if not i["dimensions"]),
            "note": "Run measure_asset (mode='batch') on items to fill in dimensions." if any(not i["dimensions"] for i in items) else None}


PAGE_CSS = """
*{box-sizing:border-box}body{font-family:Inter,Segoe UI,Arial,sans-serif;margin:0;padding:32px;color:#1d1d1f;background:#fff}
h1{font-weight:600;font-size:26px;margin:0 0 4px}h2{font-size:17px;margin:28px 0 12px;border-bottom:1px solid #ddd;padding-bottom:6px}
.sub{color:#777;font-size:13px;margin-bottom:24px}.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(220px,1fr));gap:18px}
.card{border:1px solid #e3e3e3;border-radius:8px;overflow:hidden;break-inside:avoid;background:#fff}
.card img{width:100%;aspect-ratio:4/3;object-fit:cover;background:#f2f2f2;display:block}
.noimg{width:100%;aspect-ratio:4/3;background:#f2f2f2;display:flex;align-items:center;justify-content:center;color:#aaa;font-size:12px}
.body{padding:10px 12px;font-size:12px;line-height:1.45}.name{font-weight:600;font-size:13px}.meta{color:#666}
.code{float:right;color:#999;font-size:11px}label.pick{display:flex;gap:8px;align-items:center;margin-top:8px;font-weight:600}
textarea{width:100%;margin-top:6px;font:inherit;font-size:12px;border:1px solid #ddd;border-radius:4px;padding:6px;min-height:42px}
button{margin-top:24px;padding:10px 16px;border:0;border-radius:6px;background:#1d1d1f;color:#fff;font-size:14px;cursor:pointer}
pre#out{white-space:pre-wrap;background:#f6f6f6;padding:12px;border-radius:6px;font-size:12px;display:none}
@media print{body{padding:0}button,pre#out{display:none}@page{size:A3 landscape;margin:12mm}}
@media (max-width:600px){body{padding:16px}}
"""


def _img_tag(path: str | None) -> str:
    data = _thumb_bytes(path)
    if not data:
        return '<div class="noimg">no preview</div>'
    return f'<img src="data:image/jpeg;base64,{base64.b64encode(data).decode()}" alt="">'


def _card(i: dict, extra: str = "") -> str:
    e = html.escape
    lines = [f'<span class="code">{e(i.get("code", ""))}</span><div class="name">{e(i["name"])}</div>']
    meta = " · ".join(x for x in (i["brand"], i["category"]) if x)
    if meta:
        lines.append(f'<div class="meta">{e(meta)}</div>')
    if i["dimensions"]:
        lines.append(f'<div class="meta">{e(i["dimensions"])}</div>')
    if i["materials"]:
        lines.append(f'<div class="meta">{e(i["materials"])}</div>')
    return f'<div class="card">{_img_tag(i["image"])}<div class="body">{"".join(lines)}{extra}</div></div>'


@tool
def project_board(assets: list[str] | None = None, tag: str = "", scene: str = "", title: str = "Project Board",
                  group_by_category: bool = True, output: str = "") -> dict:
    """A presentation board (one self-contained HTML page, prints to A3 landscape / PDF) of the chosen
    models with previews, brand, category, size and materials, grouped by category. Items come from a
    list of assets, a tag, or the library models used by a .max scene."""
    previews = _first_previews()
    items = [_item(r, previews) for r in _rows(assets, tag, scene)]
    if not items:
        raise ConnecterError("No items found")
    groups: dict[str, list] = defaultdict(list)
    for i in sorted(items, key=lambda i: (i["tag"], i["name"].lower())):
        groups[i["category"] if group_by_category else ""].append(i)
    body = "".join((f"<h2>{html.escape(g)}</h2>" if g else "") + '<div class="grid">' + "".join(_card(i) for i in its) + "</div>"
                   for g, its in groups.items())
    path = _out(output, title, ".html")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
                    f"<title>{html.escape(title)}</title><style>{PAGE_CSS}</style></head><body><h1>{html.escape(title)}</h1>"
                    f"<div class='sub'>{len(items)} items · {time.strftime('%d %B %Y')}</div>{body}</body></html>", "utf-8")
    return {"html": str(path), "items": len(items), "note": "Open it in a browser and print to PDF for a board."}


@tool
def client_selection_sheet(options: dict[str, list[str]], title: str = "Selections", project: str = "", output: str = "") -> dict:
    """A page for a client to choose between options: options is {'Living room sofa': [asset, asset, ...],
    'Dining chairs': [...]}. The client ticks one option per group, adds notes, and presses
    'Copy my selections' to paste their choices back to you. One self-contained HTML file to email."""
    if not options:
        raise ConnecterError("options is empty")
    previews = _first_previews()
    sections = []
    for gi, (group, assets) in enumerate(options.items()):
        cards = []
        for oi, row in enumerate(_rows(assets, "", "")):
            i = _item(row, previews)
            i["code"] = f"Option {chr(65 + oi)}"
            pick = (f'<label class="pick"><input type="radio" name="g{gi}" value="{html.escape(i["code"] + " - " + i["name"])}">'
                    f'Choose {html.escape(i["code"])}</label>')
            cards.append(_card(i, pick))
        sections.append(f'<section data-group="{html.escape(group)}"><h2>{html.escape(group)}</h2><div class="grid">{"".join(cards)}</div>'
                        f'<textarea placeholder="Notes for {html.escape(group)}"></textarea></section>')
    script = """
<script>
document.getElementById('copy').onclick = async () => {
  const lines = [document.title, ''];
  document.querySelectorAll('section').forEach((s) => {
    const pick = s.querySelector('input:checked');
    const note = s.querySelector('textarea').value.trim();
    lines.push(s.dataset.group + ': ' + (pick ? pick.value : 'no choice yet') + (note ? '\\n  Notes: ' + note : ''));
  });
  const text = lines.join('\\n');
  const out = document.getElementById('out');
  out.textContent = text; out.style.display = 'block';
  try { await navigator.clipboard.writeText(text); document.getElementById('copy').textContent = 'Copied - paste it into your reply'; } catch (e) {}
};
</script>"""
    path = _out(output, title, ".html")
    path.parent.mkdir(parents=True, exist_ok=True)
    head = f"{html.escape(project)} · " if project else ""
    path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
                    f"<title>{html.escape(title)}</title><style>{PAGE_CSS}</style></head><body><h1>{html.escape(title)}</h1>"
                    f"<div class='sub'>{head}Choose one option in each group, add any notes, then press the button below.</div>"
                    f"{''.join(sections)}<button id='copy'>Copy my selections</button><pre id='out'></pre>{script}</body></html>", "utf-8")
    return {"html": str(path), "groups": len(options), "options": sum(len(v) for v in options.values())}


@tool
def asset_usage_analytics(project_folders: list[str] | None = None, max_scenes: int = 1000, top: int = 30) -> dict:
    """Which library models get used in real projects: scans the .max scenes in project folders (without
    3ds Max; it reads the files each scene references) and reports the most used models, usage by
    category, models never used, and per-project counts. Results are cached per scene, so re-runs are fast.
    project_folders defaults to a PROJECTS folder next to the library root."""
    if not project_folders:
        with _db() as db:
            roots = [Path(r) for r in db.root_folders()]
        project_folders = [str(r.parent / "PROJECTS") for r in roots if (r.parent / "PROJECTS").exists()] or [str(r) for r in roots if "project" in str(r).lower()]
    if not project_folders:
        raise ConnecterError("Give project_folders")
    cache = json.loads(USAGE_CACHE.read_text("utf-8")) if USAGE_CACHE.exists() else {}
    scenes = []
    for folder in project_folders:
        for d, dirs, names in os.walk(folder):
            dirs[:] = [x for x in dirs if x.lower() not in ("autoback", "backup", "_backup", "old")]
            scenes += [Path(d) / n for n in names if n.lower().endswith(".max")]
            if len(scenes) >= max_scenes:
                break
    scenes = scenes[:max_scenes]
    usage: dict[str, list[str]] = {}
    for s in scenes:
        key, mtime = str(s).lower(), s.stat().st_mtime
        if key in cache and cache[key]["mtime"] == mtime:
            usage[str(s)] = cache[key]["models"]
            continue
        paths = [p for p in (_texture_status(t, s)["resolved"] for t in maxfile.max_texture_paths(s)) if p]
        models = [h["asset"] for h in library_models_for(paths)] if paths else []
        usage[str(s)] = models
        cache[key] = {"mtime": mtime, "models": models}
    USAGE_CACHE.parent.mkdir(parents=True, exist_ok=True)
    USAGE_CACHE.write_text(json.dumps(cache), "utf-8")
    counts = Counter(m for models in usage.values() for m in set(models))
    with _db() as db:
        models = [a for a in db.assets() if a.path.lower().endswith(".max")]
    tags = {a.path: (a.tags[0].split("/")[0] if a.tags else "untagged") for a in models}
    by_cat = Counter()
    for m, n in counts.items():
        by_cat[tags.get(m, "not catalogued")] += n
    per_project = Counter()
    for s, ms_ in usage.items():
        for folder in project_folders:
            if s.lower().startswith(folder.lower()):
                rel = Path(s).relative_to(folder).parts
                per_project[rel[0] if len(rel) > 1 else Path(folder).name] += len(set(ms_))
    return {"scenes_scanned": len(scenes), "scenes_using_library": sum(1 for v in usage.values() if v),
            "most_used": [{"asset": m, "scenes": n, "tag": tags.get(m)} for m, n in counts.most_common(top)],
            "usage_by_category": dict(by_cat.most_common()), "per_project": dict(per_project.most_common()),
            "library_models": len(models), "never_used": len(models) - len(counts)}
